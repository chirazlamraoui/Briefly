#!/usr/bin/env python3
"""Une slide à part : architecture 3-tiers et MVC. Ne modifie pas le deck de soutenance."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent / "slide-3tiers-mvc.pptx"

W = Inches(13.333)
H = Inches(7.5)

BG = RGBColor(0xFB, 0xF6, 0xF8)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5E, 0x4E, 0x53)
ROSE = RGBColor(0x8B, 0x4D, 0x5C)
ROSE_DEEP = RGBColor(0x6E, 0x3A, 0x48)
ROSE_SOFT = RGBColor(0xF6, 0xE8, 0xED)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE8, 0xC5, 0xCE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"


def emu(inches: float) -> Emu:
    return Inches(inches)


def no_shadow(shape):
    effect = shape._element.spPr.find(qn("a:effectLst"))
    if effect is not None:
        shape._element.spPr.remove(effect)


def box(slide, x, y, w, h, fill, line=LINE, radius=0.12):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
    if radius:
        shape.adjustments[0] = min(0.5, radius)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.25)
    no_shadow(shape)
    return shape


def oval(slide, x, y, w, h, fill, line=ROSE):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, emu(x), emu(y), emu(w), emu(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.5)
    no_shadow(shape)
    return shape


def write(shape, lines, align=PP_ALIGN.CENTER, anchor="ctr"):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = emu(0.08)
    tf.margin_right = emu(0.08)
    tf.margin_top = emu(0.04)
    tf.margin_bottom = emu(0.04)
    tf._txBody.bodyPr.set("anchor", anchor)
    for i, (text, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(1)
        run = p.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.italic = False
    return shape


def text(slide, value, x, y, w, h, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    anchor_map = {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}
    tf._txBody.bodyPr.set("anchor", anchor_map[anchor])
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = False
    return shape


def arrow(slide, x, y, w, h, points_left=False):
    kind = MSO_SHAPE.LEFT_ARROW if points_left else MSO_SHAPE.RIGHT_ARROW
    shape = slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ROSE
    shape.line.fill.background()
    no_shadow(shape)
    return shape


def both_ways(slide, x, y, w, h):
    shape = slide.shapes.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, emu(x), emu(y), emu(w), emu(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ROSE
    shape.line.fill.background()
    no_shadow(shape)
    return shape


def down_arrow(slide, x, y, w, h):
    shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, emu(x), emu(y), emu(w), emu(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ROSE
    shape.line.fill.background()
    no_shadow(shape)
    return shape


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "Architecture 3-tiers et MVC"
    prs.core_properties.author = "Chiraz Lamraoui"
    prs.core_properties.subject = "Slide annexe Briefly"

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(0), emu(0), emu(13.333), emu(0.07))
    band.fill.solid()
    band.fill.fore_color.rgb = ROSE
    band.line.fill.background()
    no_shadow(band)

    text(slide, "CONCEPTION", 0.48, 0.18, 6, 0.24, 11, True, ROSE)
    text(slide, "Architecture 3-tiers et MVC", 0.48, 0.4, 12.3, 0.48, 28, True, INK)
    text(
        slide,
        "Le 3-tiers dit où chaque partie s'exécute. Le MVC dit comment le serveur traite une requête.",
        0.48,
        0.96,
        12.3,
        0.32,
        15,
        False,
        MUTED,
    )

    # --- Gauche : 3-tiers, à partir du schéma, avec les retours ---
    box(slide, 0.38, 1.42, 8.28, 5.22, CARD, LINE, 0.06)
    text(slide, "Architecture 3-tiers", 0.58, 1.54, 4.2, 0.32, 16, True, ROSE)

    text(slide, "Niveau 1 : clients", 0.55, 1.98, 2.4, 0.26, 12, True, ROSE_DEEP, PP_ALIGN.CENTER)
    text(slide, "Niveau 2 : application", 3.72, 1.98, 2.3, 0.26, 12, True, ROSE_DEEP, PP_ALIGN.CENTER)
    text(slide, "Niveau 3 : données", 6.72, 1.98, 1.75, 0.26, 12, True, ROSE_DEEP, PP_ALIGN.CENTER)

    nav = box(slide, 0.55, 2.32, 2.4, 1.22, CARD, LINE, 0.14)
    write(nav, [("Navigateur", 15, True, INK), ("pages et session", 13, False, MUTED)])

    phone = box(slide, 0.55, 3.7, 2.4, 1.22, CARD, LINE, 0.14)
    write(phone, [("Application mobile", 15, True, INK), ("écrans et jeton", 13, False, MUTED)])

    server = box(slide, 3.72, 2.55, 2.3, 2.15, ROSE_SOFT, ROSE, 0.12)
    write(
        server,
        [
            ("Serveur d'application", 14, True, INK),
            ("Laravel", 13, False, ROSE_DEEP),
            ("règles métier", 13, False, MUTED),
        ],
    )

    data = oval(slide, 6.62, 2.62, 1.9, 2.0, ROSE_SOFT, ROSE)
    write(data, [("Données", 15, True, INK), ("base relationnelle", 12, False, MUTED)])

    # Navigateur : requête vers le serveur, HTML en retour.
    arrow(slide, 3.04, 2.48, 0.6, 0.18, points_left=False)
    text(slide, "requête", 3.0, 2.28, 0.68, 0.18, 10, False, MUTED, PP_ALIGN.CENTER)
    arrow(slide, 3.04, 3.0, 0.6, 0.18, points_left=True)
    text(slide, "HTML", 3.0, 3.2, 0.68, 0.18, 10, True, ROSE, PP_ALIGN.CENTER)

    # Mobile : requête vers le serveur, JSON en retour.
    arrow(slide, 3.04, 3.88, 0.6, 0.18, points_left=False)
    text(slide, "requête", 3.0, 3.68, 0.68, 0.18, 10, False, MUTED, PP_ALIGN.CENTER)
    arrow(slide, 3.04, 4.4, 0.6, 0.18, points_left=True)
    text(slide, "JSON", 3.0, 4.6, 0.68, 0.18, 10, True, ROSE, PP_ALIGN.CENTER)

    # Serveur et base : lecture et écriture.
    both_ways(slide, 6.1, 3.15, 0.46, 0.22)
    text(slide, "lit et", 5.95, 3.4, 0.76, 0.18, 10, True, ROSE, PP_ALIGN.CENTER)
    text(slide, "écrit", 5.95, 3.56, 0.76, 0.18, 10, True, ROSE, PP_ALIGN.CENTER)

    caption = box(slide, 0.55, 5.15, 7.9, 1.25, ROSE_SOFT, None, 0.1)
    write(
        caption,
        [
            ("Les clients parlent au serveur, jamais à la base.", 14, True, INK),
            ("Le navigateur reçoit une page HTML. Le mobile reçoit du JSON.", 13, False, MUTED),
        ],
    )

    # --- Droite : MVC ---
    box(slide, 8.82, 1.42, 4.12, 5.22, CARD, LINE, 0.06)
    text(slide, "MVC", 9.02, 1.54, 3.7, 0.3, 16, True, ROSE)
    text(slide, "À l'intérieur du niveau 2", 9.02, 1.86, 3.7, 0.26, 12, False, MUTED)

    ctrl = box(slide, 9.02, 2.24, 3.72, 1.05, CARD, LINE, 0.14)
    write(ctrl, [("Contrôleur", 16, True, INK), ("reçoit la requête", 13, False, MUTED)])

    down_arrow(slide, 9.55, 3.38, 0.26, 0.32)
    down_arrow(slide, 11.95, 3.38, 0.26, 0.32)

    model = box(slide, 9.02, 3.8, 1.78, 1.35, CARD, LINE, 0.14)
    write(model, [("Modèle", 14, True, INK), ("Eloquent", 13, True, ROSE_DEEP), ("représente les données", 12, False, MUTED)])

    view = box(slide, 10.96, 3.8, 1.78, 1.35, CARD, LINE, 0.14)
    write(view, [("Vue Blade", 14, True, INK), ("affiche la page", 12, False, MUTED), ("pour le navigateur", 12, False, MUTED)])

    mobile = box(slide, 9.02, 5.32, 3.72, 1.08, ROSE_SOFT, None, 0.1)
    write(
        mobile,
        [
            ("Mobile", 14, True, ROSE_DEEP),
            ("Même contrôleur, mêmes modèles.", 12, False, INK),
            ("JSON à la place de la page HTML.", 12, False, INK),
        ],
    )

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(0.48), emu(6.82), emu(12.38), emu(0.012))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()
    no_shadow(line)
    text(slide, "Briefly", 0.48, 6.92, 1.4, 0.28, 12, True, ROSE)
    text(slide, "Chiraz Lamraoui", 1.9, 6.92, 4, 0.28, 12, False, MUTED)
    text(slide, "C6  ·  Définir l'architecture logicielle", 6.3, 6.92, 6.5, 0.28, 12, False, ROSE, PP_ALIGN.RIGHT)

    notes = slide.notes_slide.notes_text_frame
    notes.text = (
        "Je commence par la gauche. Briefly a trois niveaux. "
        "Le niveau 1, ce sont les clients : le navigateur, avec des pages et une session, "
        "et l'application mobile, avec des écrans et un jeton. "
        "Chacun envoie une requête au serveur. "
        "Le serveur renvoie une page HTML au navigateur, et du JSON au mobile. "
        "Le niveau 2, c'est Laravel. Les règles métier sont là. "
        "Le niveau 3, c'est la base relationnelle. "
        "La flèche va dans les deux sens : Laravel lit et écrit. "
        "Les clients ne parlent jamais à la base. "
        "À droite, le MVC est à l'intérieur de ce serveur. "
        "Le contrôleur reçoit la requête. "
        "Le modèle Eloquent représente les données. "
        "La vue Blade affiche la page pour le navigateur. "
        "Le mobile passe par le même contrôleur et les mêmes modèles, "
        "et il reçoit du JSON à la place de la page HTML."
    )

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
